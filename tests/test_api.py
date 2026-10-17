"""Test suite for the docx.api module."""

import io
from pathlib import Path

import pytest

from docx.api import Document as DocumentFactoryFn
from docx.document import Document as DocumentCls
from docx.opc.constants import CONTENT_TYPE as CT

from .unitutil.file import test_file
from .unitutil.mock import FixtureRequest, Mock, class_mock, function_mock, instance_mock


class DescribeDocument:
    """Unit-test suite for `docx.api.Document` factory function."""

    @pytest.mark.parametrize("in_run", [False, True])
    def it_can_add_a_picture_from_a_path(self, tmp_path: Path, in_run: bool):
        image_path = tmp_path / "图片.png"
        image_path.write_bytes(Path(test_file("python-icon.png")).read_bytes())
        document = DocumentFactoryFn()
        target = document.add_paragraph().add_run() if in_run else document

        target.add_picture(image_path)

        assert len(document.inline_shapes) == 1
        stream = io.BytesIO()
        document.save(stream)
        stream.seek(0)
        reloaded = DocumentFactoryFn(stream)
        assert len(reloaded.inline_shapes) == 1
        package = reloaded.part.package
        assert package is not None
        (image_part,) = package.image_parts
        assert image_part.blob == image_path.read_bytes()

    def it_opens_a_docx_file(self, Package_: Mock, document_: Mock):
        document_part = Package_.open.return_value.main_document_part
        document_part.document = document_
        document_part.content_type = CT.WML_DOCUMENT_MAIN

        document = DocumentFactoryFn("foobar.docx")

        Package_.open.assert_called_once_with("foobar.docx")
        assert document is document_

    def it_opens_the_default_docx_if_none_specified(
        self, _default_docx_path_: Mock, Package_: Mock, document_: Mock
    ):
        _default_docx_path_.return_value = "default-document.docx"
        document_part = Package_.open.return_value.main_document_part
        document_part.document = document_
        document_part.content_type = CT.WML_DOCUMENT_MAIN

        document = DocumentFactoryFn()

        Package_.open.assert_called_once_with("default-document.docx")
        assert document is document_

    def it_raises_on_not_a_Word_file(self, Package_: Mock):
        Package_.open.return_value.main_document_part.content_type = "BOGUS"

        with pytest.raises(ValueError, match="file 'foobar.xlsx' is not a Word file,"):
            DocumentFactoryFn("foobar.xlsx")

    # -- fixtures --------------------------------------------------------------------------------

    @pytest.fixture
    def _default_docx_path_(self, request: FixtureRequest):
        return function_mock(request, "docx.api._default_docx_path")

    @pytest.fixture
    def document_(self, request: FixtureRequest):
        return instance_mock(request, DocumentCls)

    @pytest.fixture
    def Package_(self, request: FixtureRequest):
        return class_mock(request, "docx.api.Package")
