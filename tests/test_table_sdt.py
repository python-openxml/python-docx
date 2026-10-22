# pyright: reportPrivateUsage=false

"""Reading table cells enclosed in structured document tags (content controls)."""

from __future__ import annotations

from io import BytesIO

import pytest
from lxml import etree

from docx import Document
from docx.document import Document as DocumentObject
from docx.exceptions import InvalidSpanError
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.table import CT_Tc
from docx.table import Table


def _document(values: list[list[str]]) -> tuple[DocumentObject, Table]:
    document = Document()
    table = document.add_table(rows=len(values), cols=len(values[0]))
    for row, cells in zip(table.rows, values):
        for cell, text in zip(row.cells, cells):
            cell.text = text
    return document, table


def _wrap(tc: CT_Tc, depth: int = 1) -> None:
    parent = tc.getparent()
    assert parent is not None
    position = parent.index(tc)
    content: etree._Element = tc
    for _ in range(depth):
        sdt = OxmlElement("w:sdt")
        properties = OxmlElement("w:sdtPr")
        properties.append(OxmlElement("w:text"))
        body = OxmlElement("w:sdtContent")
        body.append(content)
        sdt.extend([properties, body])
        content = sdt
    parent.insert(position, content)


def _reopen(document: DocumentObject) -> DocumentObject:
    output = BytesIO()
    document.save(output)
    return Document(BytesIO(output.getvalue()))


def _assert_matrix(document: DocumentObject, expected: list[list[str]]) -> Table:
    document = _reopen(document)
    before = document._element.xml
    table = document.tables[0]
    assert [[cell.text for cell in row.cells] for row in table.rows] == expected
    assert [
        [table.cell(r, c).text for c in range(len(expected[0]))] for r in range(len(expected))
    ] == expected
    assert [[cell.text for cell in column.cells] for column in table.columns] == [
        list(column) for column in zip(*expected)
    ]
    for tr in table._tbl.tr_lst:
        offset = tr.grid_before
        for tc in tr.iter_tcs():
            assert tc.grid_offset == offset
            assert tr.tc_at_grid_offset(offset) is tc
            offset += tc.grid_span
    assert document._element.xml == before
    assert _reopen(document)._element.xml == before
    return table


class DescribeTableCellContentControls:
    @pytest.mark.parametrize("position", [0, 1, 2])
    @pytest.mark.parametrize("depth", [1, 2])
    def it_reads_wrapped_cells_in_grid_order_after_save(self, position: int, depth: int) -> None:
        values = [["Left", "Apple", "Right"], ["Next", "Orange", "End"]]
        document, table = _document(values)
        _wrap(table.cell(0, position)._tc, depth)
        _assert_matrix(document, values)

    def it_reads_multiple_content_controls_without_exposing_nested_table_cells(self) -> None:
        document, table = _document([["Left", "Apple", "Right"]])
        middle = table.cell(0, 1)
        nested = middle.add_table(rows=1, cols=2)
        nested.cell(0, 0).text = "Nested A"
        nested.cell(0, 1).text = "Nested B"
        # Capture references before wrapping, including a nested control.
        first, middle_tc, last = table._tbl.tr_lst[0].tc_lst
        _wrap(first)
        _wrap(middle_tc, 2)
        _wrap(last)
        reopened = _assert_matrix(document, [["Left", "Apple\n", "Right"]])
        assert len(list(reopened._tbl.iter_tcs())) == 3
        assert [cell.text for cell in reopened.cell(0, 1).tables[0].rows[0].cells] == [
            "Nested A",
            "Nested B",
        ]

    @pytest.mark.parametrize("position", [0, 1, 2])
    def it_preserves_horizontal_spans_before_inside_and_after_controls(self, position: int) -> None:
        document, table = _document([["Left", "", "", "Right"]])
        table.cell(0, 1).merge(table.cell(0, 2)).text = "Apple"
        _wrap(table._tbl.tr_lst[0].tc_lst[position])
        _assert_matrix(document, [["Left", "Apple", "Apple", "Right"]])

    @pytest.mark.parametrize(
        "wrapped_cells",
        [
            [(0, 1)],
            [(1, 1)],
            [(0, 1), (1, 1)],
            [(1, 0)],
            [(1, 2)],
        ],
    )
    def it_resolves_vertical_merges_at_the_original_grid_offset(
        self, wrapped_cells: list[tuple[int, int]]
    ) -> None:
        document, table = _document([["A0", "", "R0"], ["A1", "", "R1"]])
        table.cell(0, 1).merge(table.cell(1, 1)).text = "Apple"
        physical = [list(tr.tc_lst) for tr in table._tbl.tr_lst]
        for row, column in wrapped_cells:
            _wrap(physical[row][column])
        _assert_matrix(document, [["A0", "Apple", "R0"], ["A1", "Apple", "R1"]])

    def it_keeps_intentional_omitted_grid_positions_distinct_from_wrapped_cells(self) -> None:
        document, table = _document([["", "Apple", "Right", ""]])
        tr = table._tbl.tr_lst[0]
        first, middle, _, last = tr.tc_lst
        tr.remove(first)
        tr.remove(last)
        for tag in ("w:gridBefore", "w:gridAfter"):
            marker = OxmlElement(tag)
            marker.set(qn("w:val"), "1")
            tr.get_or_add_trPr().append(marker)
        _wrap(middle)
        row = _reopen(document).tables[0].rows[0]
        before = row._tr.xml
        assert [cell.text for cell in row.cells] == ["Apple", "Right"]
        assert row.grid_cols_before == row.grid_cols_after == 1
        assert [tc.grid_offset for tc in row._tr.iter_tcs()] == [1, 2]
        assert row._tr.tc_at_grid_offset(1) is row.cells[0]._tc
        assert row._tr.xml == before

    def it_does_not_include_deleted_cells_or_control_properties(self) -> None:
        document, table = _document([["Left", "Apple", "Right"]])
        middle = table.cell(0, 1)._tc
        _wrap(middle)
        tr = table._tbl.tr_lst[0]
        deleted = OxmlElement("w:del")
        hidden = OxmlElement("w:tc")
        deleted.append(hidden)
        tr.append(deleted)
        _assert_matrix(document, [["Left", "Apple", "Right"]])
        assert len(tr.tc_lst) == 2  # XML child access used by authoring stays unchanged.

    @pytest.mark.parametrize(
        "corners", [((0, 0), (0, 1)), ((0, 1), (0, 2)), ((0, 0), (0, 2)), ((0, 0), (1, 2))]
    )
    def it_rejects_merging_across_a_control_before_changing_any_xml(
        self, corners: tuple[tuple[int, int], tuple[int, int]]
    ) -> None:
        document, table = _document([["A0", "Apple", "C0"], ["A1", "B1", "C1"]])
        _wrap(table.cell(0, 1)._tc)
        table = _reopen(document).tables[0]
        before = table._tbl.xml
        with pytest.raises(InvalidSpanError, match="content controls"):
            table.cell(*corners[0]).merge(table.cell(*corners[1]))
        assert table._tbl.xml == before

    def it_can_merge_direct_cells_when_a_control_is_outside_the_rectangle(self) -> None:
        document, table = _document([["Apple", "Left", "Right"]])
        _wrap(table.cell(0, 0)._tc)
        control = table._tbl.xpath("./w:tr/w:sdt")[0]
        before = etree.tostring(control)
        table.cell(0, 1).merge(table.cell(0, 2))
        assert etree.tostring(control) == before
        _assert_matrix(document, [["Apple", "Left\nRight", "Left\nRight"]])
